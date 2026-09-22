# Tableau de bord YJC (Yaakaar Jeunesse Citoyennete) - Consortium Jeunesse Senegal
# Lit en direct le Google Sheet "Dashboard YJC-Deploye" : feuille Global + 5 feuilles regionales.
# Regles de lecture : ligne 1 a 3 = en-tetes ; libelle aligne a droite = desagregation ;
# libelle centre = sous-groupe (ex. Appel a projet 1) ; lignes masquees exclues par defaut.

import re
import unicodedata

SHEETS = ["Global", "Tambacounda", "Dakar", "Kedougou", "Sedhiou", "Matam"]
REGIONS = SHEETS[1:]
REGION_LABELS = {"Global": "Global (projet)", "Tambacounda": "Tambacounda", "Dakar": "Dakar",
                 "Kedougou": "Kédougou", "Sedhiou": "Sédhiou", "Matam": "Matam"}


def norm(s):
    s = str(s or "").replace("\u2019", "'").replace("\u202f", " ").replace("\xa0", " ")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"\s+", " ", s).strip().lower()
    return s.rstrip(" .:")


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _row_label(row, col):
    if col < len(row):
        v = row[col]["v"]
        if isinstance(v, str) and v.strip():
            return v.strip()
    return None


def _cell(row, col):
    return row[col] if col < len(row) else {"v": None, "pct": False, "align": None}


def detect_layout(grid):
    ind_col = None
    for r in range(min(5, len(grid))):
        for c, cell in enumerate(grid[r]):
            if isinstance(cell["v"], str) and norm(cell["v"]) == "indicateurs":
                ind_col, hdr_row = c, r
                break
        if ind_col is not None:
            break
    if ind_col is None:
        raise ValueError("Colonne 'Indicateurs' introuvable dans les 5 premières lignes")
    years_row = grid[hdr_row] if hdr_row < len(grid) else []
    q_row = grid[hdr_row + 1] if hdr_row + 1 < len(grid) else []
    year_starts = []
    for c, cell in enumerate(years_row):
        m = re.search(r"(20\d\d)", str(cell["v"] or ""))
        if m and c > ind_col:
            year_starts.append((c, int(m.group(1))))
    def year_of(col):
        y = None
        for c, yy in year_starts:
            if c <= col:
                y = yy
        return y
    quarters, annual = [], []
    for c, cell in enumerate(q_row):
        lab = str(cell["v"] or "").strip()
        if re.fullmatch(r"T\d+", lab):
            quarters.append({"q": lab, "n": int(lab[1:]), "year": year_of(c), "col": c})
        elif norm(lab).startswith("tot"):
            m = re.search(r"(20\d\d)", lab)
            annual.append({"year": int(m.group(1)) if m else year_of(c), "col": c})
    quarters.sort(key=lambda x: x["n"])
    return {"ind": ind_col, "cible": ind_col + 1, "att": ind_col + 2, "obs": ind_col + 3,
            "start": hdr_row + 3, "quarters": quarters, "annual": annual}


def classify(label, align):
    n = norm(label)
    if align == "CENTER" or n.startswith("appel a projet"):
        return "group"
    if align == "RIGHT" or n.startswith("dont") or n in {"pourcentage", "tambali", "dolele", "yaataal"}:
        return "sub"
    return "main"


def is_pct_label(label):
    n = norm(label)
    return n.startswith(("pourcentage", "%", "proportion"))


def parse_sheet(sheet, grid, hidden_rows):
    lay = detect_layout(grid)
    recs, main, group = [], None, None
    empty_run, main_count = 0, {}
    for r in range(lay["start"], len(grid)):
        row = grid[r]
        label = _row_label(row, lay["ind"])
        if not label:
            empty_run += 1
            if empty_run >= 15:
                break
            continue
        empty_run = 0
        ci, ca = _cell(row, lay["cible"]), _cell(row, lay["att"])
        kind = classify(label, _cell(row, lay["ind"])["align"])
        hidden = r in hidden_rows
        if kind == "main" or main is None:
            main_count[norm(label)] = main_count.get(norm(label), 0) + 1
            main = {"label": label, "hidden": hidden, "occ": main_count[norm(label)]}
            group, level, grp_label = None, 0, ""
        elif kind == "group":
            group = {"label": label}
            level, grp_label = 1, label
        else:
            level = 2 if group else 1
            grp_label = group["label"] if group else ""
        pct = bool(ci["pct"] or ca["pct"]) or (level == 0 and is_pct_label(label)) or norm(label) == "pourcentage"
        obs = _cell(row, lay["obs"])["v"]
        rec = {
            "sheet": sheet, "row": r + 1, "level": level, "kind": kind,
            "main": main["label"], "occ": main["occ"], "group": grp_label, "label": label,
            "cible": num(ci["v"]), "atteint": num(ca["v"]),
            "obs": obs.strip() if isinstance(obs, str) and obs.strip() else "",
            "pct": pct, "money": norm(main["label"]).startswith("montant"),
            "hidden": hidden or main["hidden"],
            "quarters": [], "annual": [],
        }
        for q in lay["quarters"]:
            rec["quarters"].append({"q": q["q"], "n": q["n"], "year": q["year"],
                                    "cible": num(_cell(row, q["col"])["v"]),
                                    "atteint": num(_cell(row, q["col"] + 1)["v"])})
        for a in lay["annual"]:
            rec["annual"].append({"year": a["year"], "cible": num(_cell(row, a["col"])["v"]),
                                  "atteint": num(_cell(row, a["col"] + 1)["v"])})
        recs.append(rec)
    return recs


def _assign_keys(recs):
    seen = {}
    for rec in recs:
        own = rec["canon"] if rec["level"] == 0 else rec["label"]
        base = f"{norm(rec['canon'])}#{rec['occ']}||{norm(rec['group'])}||{norm(own)}"
        k = (rec["sheet"], base)
        seen[k] = seen.get(k, 0) + 1
        rec["key"] = f"{base}||{seen[k]}"
        rec["main_key"] = f"{norm(rec['canon'])}#{rec['occ']}"
        c, a = rec["cible"], rec["atteint"]
        rec["taux"] = (a / c) if (c not in (None, 0) and a is not None) else None


def parse_all(grids, hidden):
    from difflib import SequenceMatcher
    out = []
    for s in SHEETS:
        if s in grids:
            out.extend(parse_sheet(s, grids[s], hidden.get(s, set())))
    ref = []
    for rec in out:
        if rec["sheet"] == "Global" and rec["level"] == 0 and rec["main"] not in ref:
            ref.append(rec["main"])
    ref_norm = {norm(x): x for x in ref}
    cache = {}
    for rec in out:
        m = rec["main"]
        if norm(m) in ref_norm or rec["sheet"] == "Global":
            rec["canon"] = m
            continue
        if m not in cache:
            best, score = m, 0.0
            for cand in ref:
                sc = SequenceMatcher(None, norm(m), norm(cand)).ratio()
                if sc > score:
                    best, score = cand, sc
            cache[m] = best if score >= 0.8 else m
        rec["canon"] = cache[m]
    _assign_keys(out)
    order = {}
    for rec in out:
        if rec["level"] == 0 and rec["main_key"] not in order:
            order[rec["main_key"]] = len(order) + 1
    for rec in out:
        rec["num"] = order.get(rec["main_key"])
    return out


def api_to_grids(payload):
    grids, hidden = {}, {}
    for sh in payload.get("sheets", []):
        title = sh["properties"]["title"]
        data = (sh.get("data") or [{}])[0]
        r0, c0 = data.get("startRow", 0), data.get("startColumn", 0)
        grid = [[] for _ in range(r0)]
        for rd in data.get("rowData", []):
            row = [{"v": None, "pct": False, "align": None} for _ in range(c0)]
            for val in rd.get("values", []) or []:
                ev = val.get("effectiveValue") or {}
                if "numberValue" in ev:
                    v = ev["numberValue"]
                elif "stringValue" in ev:
                    v = ev["stringValue"]
                elif "boolValue" in ev:
                    v = ev["boolValue"]
                else:
                    v = None
                fmt = val.get("effectiveFormat") or {}
                row.append({"v": v,
                            "pct": (fmt.get("numberFormat") or {}).get("type") == "PERCENT",
                            "align": fmt.get("horizontalAlignment")})
            grid.append(row)
        hid = set()
        for i, md in enumerate(data.get("rowMetadata", []) or []):
            if md.get("hiddenByUser"):
                hid.add(r0 + i)
        grids[title], hidden[title] = grid, hid
    return grids, hidden


def xlsx_to_grids(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    grids, hidden = {}, {}
    for s in SHEETS:
        ws = wb[s]
        grid = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 300)):
            grid.append([{"v": c.value if not (isinstance(c.value, str) and c.value.startswith("#")) else None,
                          "pct": "%" in (c.number_format or ""),
                          "align": (c.alignment.horizontal or "").upper() or None} for c in row])
        grids[s] = grid
        hidden[s] = {r - 1 for r, d in ws.row_dimensions.items() if d.hidden}
    return grids, hidden

# ============================================================
#  INTERFACE
# ============================================================
import os
import html
from datetime import datetime, timezone, timedelta

import pandas as pd
import plotly.graph_objects as go
import streamlit as st

SPREADSHEET_ID = "1i-CBHtOvvTIZPpaEMelbumbV7il-ha4utJMAmq_FvvA"
C_OK, C_MID, C_LOW, C_NONE = "#1F7A5C", "#C98A1B", "#B23A48", "#8A94A6"
C_BLUE, C_TARGET = "#2F5D8A", "#9FB3C8"
DAKAR_TZ = timezone(timedelta(hours=0))

st.set_page_config(page_title="Tableau de bord YJC", page_icon="📊", layout="wide",
                   initial_sidebar_state="expanded")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Public+Sans:wght@400;500;600;700&display=swap');
html, body, [class*="css"], .stMarkdown, .stDataFrame, button, input, select, textarea {
  font-family: 'Public Sans', 'Segoe UI', Roboto, sans-serif !important;
}
.block-container { padding-top: 2.2rem; max-width: 1400px; }
.yjc-head h1 { font-size: 1.9rem; font-weight: 700; margin: 0 0 .2rem 0; letter-spacing: -.01em; }
.yjc-head p { margin: 0; opacity: .75; font-size: .98rem; }
.yjc-stats { display: grid; grid-template-columns: repeat(4, minmax(0,1fr)); gap: 14px; margin: 1.2rem 0 1.4rem; }
.yjc-stat { border: 1px solid rgba(127,127,127,.22); border-radius: 10px; padding: 14px 16px; }
.yjc-stat .v { font-size: 2rem; font-weight: 700; line-height: 1.1; font-variant-numeric: tabular-nums; }
.yjc-stat .l { font-size: .88rem; opacity: .75; margin-top: 2px; }
.yjc-grid { display: grid; grid-template-columns: repeat(3, minmax(0,1fr)); gap: 14px; }
.yjc-card { border: 1px solid rgba(127,127,127,.22); border-left: 5px solid var(--st); border-radius: 0 10px 10px 0;
            padding: 12px 14px 12px 14px; background: rgba(127,127,127,.04); display: flex; flex-direction: column; gap: 8px; }
.yjc-card .top { display: flex; justify-content: space-between; align-items: baseline; gap: 8px; font-size: .78rem; }
.yjc-card .num { opacity: .6; font-variant-numeric: tabular-nums; }
.yjc-card .badge { color: var(--st); font-weight: 600; white-space: nowrap; }
.yjc-card .lab { font-size: .92rem; line-height: 1.35; min-height: 3.7em;
                 display: -webkit-box; -webkit-line-clamp: 3; -webkit-box-orient: vertical; overflow: hidden; }
.yjc-card .vals { display: flex; justify-content: space-between; align-items: baseline; font-variant-numeric: tabular-nums; }
.yjc-card .att { font-size: 1.45rem; font-weight: 700; }
.yjc-card .cib { font-size: .85rem; opacity: .7; }
.yjc-card .pct { font-size: 1.05rem; font-weight: 700; color: var(--st); }
.yjc-bar { position: relative; height: 8px; border-radius: 4px; background: rgba(127,127,127,.18); overflow: visible; }
.yjc-bar .fill { height: 100%; border-radius: 4px; background: var(--st); }
.yjc-bar .tick { position: absolute; top: -3px; width: 2px; height: 14px; background: rgba(127,127,127,.7); }
.yjc-note { font-size: .85rem; opacity: .7; margin: .3rem 0 .8rem; }
.yjc-legend { display: flex; gap: 18px; flex-wrap: wrap; font-size: .85rem; margin: .2rem 0 1rem; }
.yjc-legend span::before { content: ""; display: inline-block; width: 10px; height: 10px; border-radius: 2px;
                           background: var(--c); margin-right: 6px; vertical-align: -1px; }
@media (max-width: 1100px) { .yjc-grid { grid-template-columns: repeat(2, minmax(0,1fr)); } }
@media (max-width: 700px) { .yjc-grid, .yjc-stats { grid-template-columns: minmax(0,1fr); } }
</style>
""", unsafe_allow_html=True)


# ---------- chargement ----------
FIELDS = ("sheets(properties(title),data(startRow,startColumn,rowMetadata(hiddenByUser),"
          "rowData(values(effectiveValue,effectiveFormat(horizontalAlignment,numberFormat(type))))))")


@st.cache_data(ttl=900, show_spinner="Lecture du Google Sheet...")
def load_records():
    local = os.environ.get("YJC_LOCAL_XLSX")
    if local:
        grids, hidden = xlsx_to_grids(local)
    else:
        from google.oauth2.service_account import Credentials
        from google.auth.transport.requests import AuthorizedSession
        creds = Credentials.from_service_account_info(
            dict(st.secrets["google_credentials"]),
            scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
        session = AuthorizedSession(creds)
        resp = session.get(
            f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}",
            params={"includeGridData": "true", "fields": FIELDS,
                    "ranges": [f"'{s}'!A1:BD400" for s in SHEETS]},
            timeout=60)
        if resp.status_code == 403:
            raise PermissionError("Accès refusé au Google Sheet. Partage-le en lecture avec l'adresse "
                                  "client_email du compte de service.")
        resp.raise_for_status()
        grids, hidden = api_to_grids(resp.json())
    recs = parse_all(grids, hidden)
    return recs, datetime.now(timezone.utc).isoformat()


# ---------- formats ----------
def fmt_n(x, dec=0):
    if x is None:
        return "n.d."
    if abs(x - round(x)) < 1e-9:
        s = f"{int(round(x)):,}"
    else:
        s = f"{x:,.{dec or 1}f}"
    return s.replace(",", "\u202f").replace(".", ",")


def fmt_v(rec, x):
    if x is None:
        return "n.d."
    if rec["pct"]:
        return f"{fmt_n(x * 100, 1 if abs(x * 100 - round(x * 100)) > 0.05 else 0)}\u202f%"
    if rec["money"]:
        return f"{fmt_n(x / 1e6, 1)}\u202fM FCFA" if abs(x) >= 1e6 else f"{fmt_n(x)}\u202fFCFA"
    return fmt_n(x, 1)


def fmt_pct(t):
    if t is None:
        return "n.d."
    p = t * 100
    if p >= 100 or abs(p - round(p)) < 0.05:
        return f"{fmt_n(float(round(p)))}\u202f%"
    return f"{fmt_n(p, 1)}\u202f%"


def status(t):
    if t is None:
        return "Sans cible", C_NONE
    if t >= 1:
        return "Cible atteinte", C_OK
    if t >= 0.5:
        return "En progression", C_MID
    return "À renforcer", C_LOW


def short(s, n=90):
    s = " ".join(str(s).split())
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


def plot(fig, h=None):
    fig.update_layout(margin=dict(l=10, r=10, t=40, b=10), font=dict(family="Public Sans, sans-serif", size=13),
                      legend=dict(orientation="h", y=1.08, x=0), hoverlabel=dict(font_family="Public Sans"))
    if h:
        fig.update_layout(height=h)
    st.plotly_chart(fig, width="stretch", config={"displaylogo": False})


# ---------- données ----------
try:
    RECS, LOADED_AT = load_records()
except Exception as e:
    st.error(f"Impossible de lire le Google Sheet : {e}")
    st.info("Vérifie les secrets Streamlit (bloc [google_credentials]) et le partage du Sheet avec le compte de service.")
    st.stop()

with st.sidebar:
    st.markdown("### Paramètres")
    _scope_lab = st.selectbox("Périmètre", [REGION_LABELS[s] for s in SHEETS], key="scope",
                              help="Global regroupe toutes les régions et les activités nationales.")
    scope = next(s for s in SHEETS if REGION_LABELS[s] == _scope_lab)
    show_hidden = st.toggle("Inclure les lignes masquées", value=False,
                            help="Par défaut, les lignes masquées dans le Google Sheet sont exclues.")
    st.divider()
    if st.button("Actualiser les données", width="stretch"):
        st.cache_data.clear()
        st.rerun()
    loaded = datetime.fromisoformat(LOADED_AT).astimezone(DAKAR_TZ)
    st.caption(f"Données lues le {loaded:%d/%m/%Y à %H:%M} (heure de Dakar). "
               "Relecture automatique toutes les 15 minutes.")
    st.divider()
    st.markdown(f'<div class="yjc-legend" style="flex-direction:column;gap:6px">'
                f'<span style="--c:{C_OK}">Cible atteinte (100 % et plus)</span>'
                f'<span style="--c:{C_MID}">En progression (50 à 99 %)</span>'
                f'<span style="--c:{C_LOW}">À renforcer (moins de 50 %)</span>'
                f'<span style="--c:{C_NONE}">Sans cible définie</span></div>', unsafe_allow_html=True)

VIS = [r for r in RECS if show_hidden or not r["hidden"]]
by_sheet = {s: [r for r in VIS if r["sheet"] == s] for s in SHEETS}
mains = [r for r in by_sheet[scope] if r["level"] == 0]
ref_mains = {r["main_key"]: r for r in RECS if r["level"] == 0 and r["sheet"] == "Global"}


def main_label(mk):
    r = ref_mains.get(mk) or next((x for x in RECS if x["main_key"] == mk and x["level"] == 0), None)
    return f"{r['num']}. {short(r['label'], 110)}" if r else mk


def pick_box(label, keys, fmt, key):
    labels, seen = {}, {}
    for k in keys:
        l = fmt(k)
        seen[l] = seen.get(l, 0) + 1
        if seen[l] > 1:
            l = f"{l} ({seen[l]})"
        labels[l] = k
    if not labels:
        return None
    chosen = st.selectbox(label, list(labels), key=key)
    return labels[chosen]


st.markdown(f"""<div class="yjc-head"><h1>Yaakaar Jeunesse Citoyenneté</h1>
<p>Suivi des indicateurs du projet, périmètre : <b>{REGION_LABELS[scope]}</b></p></div>""",
            unsafe_allow_html=True)

tab1, tab2, tab3, tab4, tab5 = st.tabs(["Vue d'ensemble", "Comparaison régionale", "Évolution trimestrielle",
                                        "Fiche indicateur", "Données"])

# ============ 1. VUE D'ENSEMBLE ============
with tab1:
    with_target = [r for r in mains if r["taux"] is not None]
    n_ok = sum(r["taux"] >= 1 for r in with_target)
    n_mid = sum(0.5 <= r["taux"] < 1 for r in with_target)
    n_low = sum(r["taux"] < 0.5 for r in with_target)
    st.markdown(f"""<div class="yjc-stats">
      <div class="yjc-stat"><div class="v">{len(mains)}</div><div class="l">indicateurs suivis</div></div>
      <div class="yjc-stat"><div class="v" style="color:{C_OK}">{n_ok}</div><div class="l">cibles atteintes ou dépassées</div></div>
      <div class="yjc-stat"><div class="v" style="color:{C_MID}">{n_mid}</div><div class="l">en progression (50 à 99 %)</div></div>
      <div class="yjc-stat"><div class="v" style="color:{C_LOW}">{n_low}</div><div class="l">à renforcer (moins de 50 %)</div></div>
    </div>""", unsafe_allow_html=True)

    c1, c2, c3 = st.columns([2, 2, 1])
    q = c1.text_input("Rechercher un indicateur", placeholder="débats, cohorte, centres, Yaakaar...")
    stat_f = c2.multiselect("Statut", ["Cible atteinte", "En progression", "À renforcer", "Sans cible"],
                            placeholder="Tous les statuts")
    sort_by = c3.selectbox("Trier par", ["Ordre du Sheet", "Taux croissant", "Taux décroissant"])

    shown = [r for r in mains if (not q or norm(q) in norm(r["label"]))
             and (not stat_f or status(r["taux"])[0] in stat_f)]
    if sort_by != "Ordre du Sheet":
        shown.sort(key=lambda r: (r["taux"] is None, r["taux"] or 0), reverse=(sort_by == "Taux décroissant"))
        if sort_by == "Taux décroissant":
            shown.sort(key=lambda r: r["taux"] is None)

    if not shown:
        st.info("Aucun indicateur ne correspond à ces filtres. Modifie la recherche ou le statut.")
    cards = []
    for r in shown:
        lab, col = status(r["taux"])
        t = r["taux"] or 0
        scale = max(1.0, min(t, 2.0))
        fill = min(t, scale) / scale * 100
        tick = 100 / scale
        cards.append(f"""<div class="yjc-card" style="--st:{col}">
          <div class="top"><span class="num">Indicateur {r['num']}</span><span class="badge">{lab}</span></div>
          <div class="lab" title="{html.escape(r['label'])}">{html.escape(r['label'])}</div>
          <div class="vals"><span><span class="att">{fmt_v(r, r['atteint'])}</span>
            <span class="cib"> / {fmt_v(r, r['cible'])}</span></span><span class="pct">{fmt_pct(r['taux'])}</span></div>
          <div class="yjc-bar"><div class="fill" style="width:{fill:.1f}%"></div>
            <div class="tick" style="left:calc({tick:.1f}% - 1px)"></div></div>
        </div>""")
    st.markdown(f'<div class="yjc-grid">{"".join(cards)}</div>', unsafe_allow_html=True)
    st.markdown('<p class="yjc-note">Le trait vertical sur chaque barre marque la cible. '
                'Au-delà de 100 %, la barre est remise à l\'échelle pour montrer le dépassement.</p>',
                unsafe_allow_html=True)

    st.markdown("#### Désagrégations")
    subs = [r for r in by_sheet[scope] if r["level"] > 0]
    opts = ["Tous les indicateurs"] + [m["main_key"] for m in mains if any(s["main_key"] == m["main_key"] for s in subs)]
    pick = pick_box("Indicateur", opts, lambda k: k if k == "Tous les indicateurs" else main_label(k), "desag_pick")
    rows = []
    for r in by_sheet[scope]:
        if pick != "Tous les indicateurs" and r["main_key"] != pick:
            continue
        if pick == "Tous les indicateurs" and r["level"] == 0 and not any(s["main_key"] == r["main_key"] for s in subs):
            continue
        rows.append({"N°": r["num"],
                     "Niveau": "Indicateur" if r["level"] == 0 else ("Sous-groupe" if r["kind"] == "group" else "Désagrégation"),
                     "Libellé": ("" if r["level"] == 0 else ("    " if r["level"] == 1 else "        ")) + r["label"],
                     "Cible": fmt_v(r, r["cible"]), "Atteint": fmt_v(r, r["atteint"]),
                     "Taux": r["taux"] * 100 if r["taux"] is not None else None,
                     "Statut": status(r["taux"])[0]})
    if rows:
        st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch",
                     height=min(38 * len(rows) + 40, 620),
                     column_config={"Taux": st.column_config.ProgressColumn("Taux", format="%.0f %%",
                                                                            min_value=0, max_value=100),
                                    "Libellé": st.column_config.TextColumn(width="large")})
    else:
        st.info("Cet indicateur n'a pas de désagrégation.")

# ============ 2. COMPARAISON RÉGIONALE ============
with tab2:
    st.markdown("Taux de réalisation de chaque indicateur par région (valeur atteinte rapportée à la cible régionale).")
    reg_mains = []
    for mk, gm in ref_mains.items():
        if not show_hidden and gm["hidden"]:
            continue
        if any(r["main_key"] == mk and r["level"] == 0 for s in REGIONS for r in by_sheet[s]):
            reg_mains.append(mk)
    if not reg_mains:
        st.info("Aucun indicateur n'est suivi au niveau régional.")
    else:
        idx = {(r["sheet"], r["key"]): r for r in VIS}
        z, txt, ylab, hm_keys = [], [], [], []
        for mk in reg_mains:
            gm = ref_mains[mk]
            zr, tr = [], []
            for s in REGIONS:
                r = idx.get((s, gm["key"]))
                t = r["taux"] if r else None
                zr.append(None if t is None else (2 if t >= 1 else (1 if t >= 0.5 else 0)))
                tr.append("" if t is None else fmt_pct(t))
            if all(v is None for v in zr):
                continue
            ylab.append(f"{gm['num']}. {short(gm['label'], 62)}")
            z.append(zr)
            txt.append(tr)
            hm_keys.append(mk)
        no_target = [ref_mains[mk] for mk in reg_mains if mk not in hm_keys]
        fig = go.Figure(go.Heatmap(
            z=z, x=[REGION_LABELS[s] for s in REGIONS], y=ylab, text=txt, texttemplate="%{text}",
            zmin=0, zmax=2, xgap=3, ygap=3, showscale=False,
            colorscale=[[0, C_LOW], [0.3333, C_LOW], [0.3334, C_MID], [0.6666, C_MID], [0.6667, C_OK], [1, C_OK]],
            textfont=dict(color="white", size=13),
            hovertemplate="%{y}<br>%{x} : %{text}<extra></extra>"))
        fig.update_yaxes(autorange="reversed", tickfont=dict(size=12))
        fig.update_xaxes(side="top")
        plot(fig, h=max(420, 32 * len(ylab) + 80))
        if no_target:
            st.caption("Sans cible régionale renseignée, donc absents de la carte : "
                       + " ; ".join(f"{m['num']}. {short(m['label'], 60)}" for m in no_target))

        st.markdown("#### Détail d'un indicateur par région")
        mk = pick_box("Indicateur", reg_mains, main_label, "reg_pick")
        gkeys = [r for r in by_sheet["Global"] if r["main_key"] == mk]
        sub_opts = [r["key"] for r in gkeys]
        sub_lab = {r["key"]: ("Total de l'indicateur" if r["level"] == 0 else
                              (f"{r['group']} > {r['label']}" if r["level"] == 2 else r["label"])) for r in gkeys}
        sk = pick_box("Niveau de lecture", sub_opts, lambda k: sub_lab[k], "reg_sub")
        data = []
        for s in REGIONS:
            r = idx.get((s, sk))
            if r:
                data.append({"Région": REGION_LABELS[s], "Cible": r["cible"], "Atteint": r["atteint"],
                             "Taux": r["taux"], "rec": r})
        if data:
            ref = data[0]["rec"]
            fig = go.Figure()
            fig.add_bar(x=[d["Région"] for d in data], y=[d["Cible"] for d in data], name="Cible",
                        marker_color=C_TARGET, hovertemplate="%{x}<br>Cible : %{y:,.0f}<extra></extra>")
            fig.add_bar(x=[d["Région"] for d in data], y=[d["Atteint"] for d in data], name="Atteint",
                        marker_color=[status(d["Taux"])[1] for d in data],
                        text=[fmt_pct(d["Taux"]) for d in data], textposition="outside",
                        hovertemplate="%{x}<br>Atteint : %{y:,.0f}<extra></extra>")
            fig.update_layout(barmode="group", bargap=0.3, separators=", ")
            if ref["pct"]:
                fig.update_yaxes(tickformat=".0%")
            plot(fig, h=420)
            st.dataframe(pd.DataFrame([{"Région": d["Région"], "Cible": fmt_v(d["rec"], d["Cible"]),
                                        "Atteint": fmt_v(d["rec"], d["Atteint"]), "Taux": fmt_pct(d["Taux"]),
                                        "Statut": status(d["Taux"])[0]} for d in data]),
                         hide_index=True, width="stretch")
        else:
            st.info("Ce niveau de lecture n'existe pas dans les feuilles régionales.")

# ============ 3. ÉVOLUTION TRIMESTRIELLE ============
with tab3:
    def has_q(r):
        return any((q["atteint"] or 0) != 0 for q in r["quarters"])
    q_recs = [r for r in by_sheet[scope] if has_q(r)]
    if not q_recs:
        st.info(f"Aucune donnée trimestrielle n'est encore saisie pour {REGION_LABELS[scope]}.")
    else:
        st.markdown(f"{len({r['main_key'] for r in q_recs})} indicateurs disposent d'une saisie trimestrielle "
                    f"pour {REGION_LABELS[scope]}. Les autres sont renseignés uniquement en valeur cumulée.")
        q_mains = list(dict.fromkeys(r["main_key"] for r in q_recs))
        mk = pick_box("Indicateur", q_mains, main_label, "q_pick")
        lines = [r for r in q_recs if r["main_key"] == mk]
        lk = pick_box("Niveau de lecture", [r["key"] for r in lines],
                      lambda k: next(("Total de l'indicateur" if r["level"] == 0 else
                                      (f"{r['group']} > {r['label']}" if r["level"] == 2 else r["label"]))
                                     for r in lines if r["key"] == k), "q_sub")
        r = next(x for x in lines if x["key"] == lk)
        qs = r["quarters"]
        last = max((i for i, q in enumerate(qs) if (q["atteint"] or 0) != 0), default=len(qs) - 1)
        qs = qs[: max(last + 1, 4)]
        x = [f"{q['q']}<br>{q['year']}" for q in qs]
        att = [q["atteint"] or 0 for q in qs]
        cum, s_ = [], 0
        for a in att:
            s_ += a
            cum.append(s_)
        fig = go.Figure()
        fig.add_bar(x=x, y=att, name="Atteint sur le trimestre", marker_color=C_BLUE,
                    hovertemplate="%{x}<br>Atteint : %{y:,.0f}<extra></extra>")
        if any(q["cible"] for q in qs):
            fig.add_scatter(x=x, y=[q["cible"] for q in qs], name="Cible du trimestre", mode="lines+markers",
                            line=dict(color=C_MID, dash="dot", width=2))
        if not r["pct"]:
            fig.add_scatter(x=x, y=cum, name="Cumul depuis T1", mode="lines+markers", yaxis="y2",
                            line=dict(color=C_OK, width=3))
            fig.update_layout(yaxis2=dict(overlaying="y", side="right", showgrid=False, title="Cumul",
                                          rangemode="tozero", tickformat=",.0f"))
        fig.update_layout(separators=", ", yaxis=dict(rangemode="tozero", tickformat=",.0f"))
        plot(fig, h=450)
        st.caption(f"Affichage jusqu'au dernier trimestre renseigné ({qs[-1]['q']} {qs[-1]['year']}).")
        yrs = [a for a in r["annual"] if a["cible"] or a["atteint"]]
        if yrs:
            st.dataframe(pd.DataFrame([{"Année": str(a["year"]), "Cible annuelle": fmt_v(r, a["cible"]),
                                        "Atteint": fmt_v(r, a["atteint"]),
                                        "Taux": fmt_pct(a["atteint"] / a["cible"]) if a["cible"] and a["atteint"] is not None else "n.d."}
                                       for a in yrs]), hide_index=True, width="stretch")

# ============ 4. FICHE INDICATEUR ============
with tab4:
    all_mk = [m["main_key"] for m in mains]
    if not all_mk:
        st.info("Aucun indicateur pour ce périmètre.")
    else:
        mk = pick_box("Indicateur", all_mk, main_label, "fiche_pick")
        m = next(x for x in mains if x["main_key"] == mk)
        lab, col = status(m["taux"])
        st.markdown(f"##### {html.escape(m['label'])}")
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("Valeur atteinte", fmt_v(m, m["atteint"]))
        k2.metric("Cible", fmt_v(m, m["cible"]))
        k3.metric("Taux de réalisation", fmt_pct(m["taux"]))
        k4.markdown(f"<div style='padding-top:1.6rem;font-weight:600;color:{col}'>{lab}</div>", unsafe_allow_html=True)
        if m["obs"]:
            st.markdown(f"**Observations :** {html.escape(m['obs'])}")
        lines = [r for r in by_sheet[scope] if r["main_key"] == mk and r["level"] > 0]
        if lines:
            st.markdown("**Désagrégations**")
            fig = go.Figure(go.Bar(
                y=[(f"{r['group']} > " if r["level"] == 2 else "") + short(r["label"], 50) for r in lines],
                x=[(r["taux"] or 0) * 100 for r in lines], orientation="h",
                marker_color=[status(r["taux"])[1] for r in lines],
                text=[f"{fmt_v(r, r['atteint'])} / {fmt_v(r, r['cible'])}" if r["cible"] is not None
                      else fmt_v(r, r["atteint"]) for r in lines],
                textposition="outside", cliponaxis=False,
                hovertemplate="%{y}<br>Taux : %{x:.0f} %<extra></extra>"))
            fig.add_vline(x=100, line_dash="dot", line_color="#888")
            fig.update_yaxes(autorange="reversed")
            fig.update_xaxes(title="Taux de réalisation (%)", range=[0, max(130, max(((r["taux"] or 0) * 100 for r in lines), default=0) * 1.25)])
            plot(fig, h=max(260, 34 * len(lines) + 90))
        if scope == "Global":
            idx = {(r["sheet"], r["key"]): r for r in VIS}
            reg_rows = [{"Région": REGION_LABELS[s], "Cible": fmt_v(m, idx[(s, m["key"])]["cible"]),
                         "Atteint": fmt_v(m, idx[(s, m["key"])]["atteint"]),
                         "Taux": fmt_pct(idx[(s, m["key"])]["taux"])}
                        for s in REGIONS if (s, m["key"]) in idx]
            if reg_rows:
                st.markdown("**Répartition régionale**")
                st.dataframe(pd.DataFrame(reg_rows), hide_index=True, width="stretch")
            else:
                st.caption("Indicateur suivi uniquement au niveau du projet (pas de ventilation régionale).")

# ============ 5. DONNÉES ============
with tab5:
    df = pd.DataFrame([{"Périmètre": REGION_LABELS[r["sheet"]], "N°": r["num"],
                        "Indicateur principal": r["main"], "Sous-groupe": r["group"], "Libellé": r["label"],
                        "Niveau": r["level"], "Cible": r["cible"], "Atteint": r["atteint"],
                        "Taux (%)": round(r["taux"] * 100, 1) if r["taux"] is not None else None,
                        "Statut": status(r["taux"])[0], "Masqué dans le Sheet": r["hidden"],
                        "Ligne du Sheet": r["row"], "Observations": r["obs"]} for r in VIS])
    st.markdown(f"{len(df)} lignes lues sur les 6 feuilles (Global et 5 régions).")
    st.dataframe(df, hide_index=True, width="stretch", height=560)
    st.download_button("Télécharger en CSV", df.to_csv(index=False, sep=";").encode("utf-8-sig"),
                       file_name=f"indicateurs_yjc_{datetime.now():%Y%m%d}.csv", mime="text/csv")
